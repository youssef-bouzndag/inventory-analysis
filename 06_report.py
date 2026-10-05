import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HOLDING_RATE = 0.20   # same assumption as Lesson 5. You can also change it inside the Excel file.

#Load everything the earlier lessons produced 
products = pd.read_csv("products.csv")
demand = pd.read_csv("demand.csv")
abc = pd.read_csv("abc.csv")[["product_id", "class"]]
policies = pd.read_csv("policy_comparison.csv", index_col=0)
safety = pd.read_csv("safety_stock.csv")[["product_id", "safety_stock"]]

units = demand.groupby("product_id")["units"].sum().rename("units_sold")
table = (products.merge(abc, on="product_id")
                 .merge(units, on="product_id")
                 .merge(safety, on="product_id"))
table = table.sort_values("units_sold", ascending=False).reset_index(drop=True)

weekly = (demand.merge(products[["product_id", "category"]], on="product_id")
                .groupby(["week", "category"])["units"].sum().unstack().reset_index())

#Styles
FONT = "Arial"
head_font = Font(name=FONT, bold=True, color="FFFFFF")
head_fill = PatternFill("solid", fgColor="1F4E78")
input_font = Font(name=FONT, color="0000FF", bold=True)       # blue = input you can change
input_fill = PatternFill("solid", fgColor="FFFF00")
base = Font(name=FONT)
bold = Font(name=FONT, bold=True)


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font, cell.fill = head_font, head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def autowidth(ws, minimum=11, maximum=46):
    for col in ws.columns:
        longest = max((len(str(c.value)) for c in col if c.value is not None), default=0)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(minimum, longest + 2), maximum)


def set_font(ws):
    for row in ws.iter_rows():
        for c in row:
            if c.font == Font() or c.font.name != FONT:
                c.font = Font(name=FONT, bold=c.font.bold, color=c.font.color)


def polish(chart, x_title, y_title, y_format="#,##0", legend=None, left=0.12):
    """Axes visible + titled, and a fixed plot area so titles, labels and legend never overlap."""
    chart.x_axis.title = x_title
    chart.y_axis.title = y_title
    chart.x_axis.delete = False          # openpyxl hides axes (and their values) unless this is False
    chart.y_axis.delete = False
    chart.y_axis.number_format = y_format
    for t in (chart.title, chart.x_axis.title, chart.y_axis.title):
        t.overlay = False                # titles must not sit on top of the plot
    if legend is None:
        chart.legend = None
        plot_h = 0.58
    else:
        chart.legend.position = legend   # "b" = below the chart
        chart.legend.overlay = False
        plot_h = 0.52
    # Reserve space: top 14% for the title, `left` share for y labels/title, bottom for x labels/title/legend
    chart.layout = Layout(manualLayout=ManualLayout(
        layoutTarget="inner", xMode="edge", yMode="edge",
        x=left, y=0.14, w=0.97 - left, h=plot_h))


wb = Workbook()

#Products sheet (data + formulas)
wp = wb.active
wp.title = "Products"
cols = ["Product", "Category", "Class", "Unit cost", "Unit price", "Lead time (weeks)",
        "Units sold", "Revenue", "Profit", "Safety stock (units)"]
wp.append(cols)
for i, r in table.iterrows():
    n = i + 2
    wp.append([r["product_id"], r["category"], r["class"], r["unit_cost"], r["unit_price"],
               int(r["lead_time_weeks"]), int(r["units_sold"]),
               f"=G{n}*E{n}", f"=G{n}*(E{n}-D{n})", int(r["safety_stock"])])
last = len(table) + 1
style_header(wp, 1, len(cols))
for r in range(2, last + 1):
    for c in (4, 5):
        wp.cell(r, c).number_format = "#,##0.00"
    for c in (7, 8, 9, 10):
        wp.cell(r, c).number_format = "#,##0"
wp.freeze_panes = "A2"
wp.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last}"

#Policies sheet (holding rate is an input)
wl = wb.create_sheet("Policies")
wl["A1"], wl["B1"] = "Holding cost (share of stock value per year)", HOLDING_RATE
wl["B1"].number_format = "0%"
wl["A1"].font = bold
wl["B1"].font, wl["B1"].fill = input_font, input_fill
wl["A2"] = "Blue/yellow cell = assumption. Change it and the costs below recalculate."
wl["A2"].font = Font(name=FONT, italic=True)

heads = ["Policy", "Lost revenue", "Stock value", "Fill rate %", "Lost profit",
         "Holding cost", "Total cost"]
wl.append(heads)   # row 3 (row 2 holds the note)
style_header(wl, 3, len(heads))
for k, (name, r) in enumerate(policies.iterrows()):
    n = 4 + k
    wl.append([name, r["lost_revenue"], r["stock_value"], r["fill_rate_%"], r["lost_profit"],
               f"=C{n}*$B$1", f"=E{n}+F{n}"])
p_last = 3 + len(policies)
for r in range(4, p_last + 1):
    for c in (2, 3, 5, 6, 7):
        wl.cell(r, c).number_format = "#,##0"
    wl.cell(r, 4).number_format = "0.0"
wl.cell(p_last + 2, 1, "Lost profit = lost units x (price - cost). Holding cost = average stock value x holding rate.")
wl.cell(p_last + 2, 1).font = Font(name=FONT, italic=True)

chart = BarChart()
chart.type, chart.title = "col", "Total cost by policy (lower is better)"
chart.add_data(Reference(wl, min_col=7, min_row=3, max_row=p_last), titles_from_data=True)
chart.set_categories(Reference(wl, min_col=1, min_row=4, max_row=p_last))
polish(chart, "Policy", "Total cost (lost profit + holding cost)")
chart.height, chart.width = 10, 24
wl.add_chart(chart, f"A{p_last + 4}")

#Class sheet (formulas on the Products sheet)
wc = wb.create_sheet("By class")
wc.append(["Class", "Products", "Units sold", "Revenue", "Share of revenue", "Safety stock (units)"])
style_header(wc, 1, 6)
for k, cl in enumerate(["A", "B", "C"]):
    n = 2 + k
    wc.append([cl, f'=COUNTIF(Products!$C$2:$C${last},A{n})',
               f'=SUMIFS(Products!$G$2:$G${last},Products!$C$2:$C${last},A{n})',
               f'=SUMIFS(Products!$H$2:$H${last},Products!$C$2:$C${last},A{n})',
               f'=IFERROR(D{n}/$D$5,0)',
               f'=SUMIFS(Products!$J$2:$J${last},Products!$C$2:$C${last},A{n})'])
wc.append(["Total", "=SUM(B2:B4)", "=SUM(C2:C4)", "=SUM(D2:D4)", "=SUM(E2:E4)", "=SUM(F2:F4)"])
for c in range(1, 7):
    wc.cell(5, c).font = bold
for r in range(2, 6):
    for c in (3, 4, 6):
        wc.cell(r, c).number_format = "#,##0"
    wc.cell(r, 5).number_format = "0.0%"
ch2 = BarChart()
ch2.type, ch2.title = "col", "Revenue by ABC class"
ch2.add_data(Reference(wc, min_col=4, min_row=1, max_row=4), titles_from_data=True)
ch2.set_categories(Reference(wc, min_col=1, min_row=2, max_row=4))
polish(ch2, "ABC class", "Revenue (millions)", y_format='0.0,,"M"', left=0.17)
ch2.height, ch2.width = 9, 15
wc.add_chart(ch2, "A8")

#Weekly demand sheet + seasonality chart
ww = wb.create_sheet("Weekly demand")
ww.append(["Week"] + list(weekly.columns[1:]))
for row in weekly.itertuples(index=False):
    ww.append([int(row[0])] + [int(v) for v in row[1:]])
style_header(ww, 1, weekly.shape[1])
line = LineChart()
line.title = "Weekly units by category"
line.add_data(Reference(ww, min_col=2, max_col=weekly.shape[1], min_row=1, max_row=len(weekly) + 1),
              titles_from_data=True)
line.set_categories(Reference(ww, min_col=1, min_row=2, max_row=len(weekly) + 1))
polish(line, "Week", "Units", legend="b")
line.x_axis.tickLblSkip = 4          # show every 4th week label so numbers do not crowd
line.height, line.width = 10, 24
ww.add_chart(line, "G2")

#Summary sheet (first tab) 
ws = wb.create_sheet("Summary", 0)
ws["A1"] = "Inventory analysis: summary"
ws["A1"].font = Font(name=FONT, bold=True, size=16)
kpis = [
    ("Products", f"=COUNTA(Products!A2:A{last})", "#,##0"),
    ("Total units sold", f"=SUM(Products!G2:G{last})", "#,##0"),
    ("Total revenue", f"=SUM(Products!H2:H{last})", "#,##0"),
    ("Total profit", f"=SUM(Products!I2:I{last})", "#,##0"),
    ("Profit margin", "=B6/B5", "0.0%"),
    ("Revenue from class A products", "='By class'!E2", "0.0%"),
    ("Total cost, policy 1 (average demand)", "=Policies!G4", "#,##0"),
    ("Total cost, policy 2 (seasonal forecast)", "=Policies!G5", "#,##0"),
    ("Total cost, policy 3 (forecast + safety stock)", "=Policies!G6", "#,##0"),
    
]
for k, (label, f, fmt) in enumerate(kpis):
    r = 3 + k
    ws.cell(r, 1, label).font = base
    ws.cell(r, 2, f).number_format = fmt
    ws.cell(r, 2).font = bold
ws.column_dimensions["A"].width = 48
ws.column_dimensions["B"].width = 16

#Finish
for sheet in (wp, wl, wc, ww):
    autowidth(sheet)
wl.column_dimensions["A"].width = 48
for sheet in wb.worksheets:
    set_font(sheet)

for sheet in (wl, wc, ww, ws):           # print each chart sheet on one landscape page
    sheet.page_setup.orientation = "landscape"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1

wb.save("inventory_report.xlsx")
print("Saved inventory_report.xlsx with sheets:", wb.sheetnames)